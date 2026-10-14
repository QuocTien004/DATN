import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

md_path = Path('d:/DoAnTotNghiep/BC_TienDo/12-10_18-10(W7).md')
docx_path = Path('d:/DoAnTotNghiep/BC_TienDo/12-10_18-10(W7).docx')

md_content = md_path.read_text(encoding='utf-8')

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
COLOR_PRIMARY = RGBColor(26, 54, 93)     # Navy #1A365D
COLOR_SECONDARY = RGBColor(43, 108, 176) # Blue #2B6CB0
COLOR_TEXT = RGBColor(45, 55, 72)        # Charcoal #2D3748

def set_para_font(p, name='Times New Roman', size=12, color=COLOR_TEXT, bold=False, italic=False, space_after=6):
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    for r in p.runs:
        r.font.name = name
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        r.italic = italic

# Title
p_title = doc.add_paragraph()
r_title = p_title.add_run('BÁO CÁO TIẾN ĐỘ ĐỒ ÁN TỐT NGHIỆP')
set_para_font(p_title, size=16, color=COLOR_PRIMARY, bold=True, space_after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_sub = doc.add_paragraph()
p_sub.add_run('Tuần 7 (12/10/2026 – 18/10/2026)\nĐề tài: Mô hình thế giới (World Model) cho xe tự lái trong môi trường MetaDrive\nNgười thực hiện: Nhánh Bao (worldmodel-metadrive_Bao)')
set_para_font(p_sub, size=11, color=COLOR_TEXT, italic=True, space_after=12)
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

def sanitize_text(text: str) -> str:
    # Remove markdown formatting
    text = text.replace('**', '').replace('`', '').replace('*', '')
    text = text.replace('$\\Delta\\text{lat}$', 'Δlat')
    text = text.replace('$\\frac{\\partial J}{\\partial \\text{throttle}}$', '∂J/∂throttle')
    text = text.replace('$', '')
    # Remove any illegal XML control characters (ASCII 0-31 except 9, 10, 13)
    clean_chars = [c for c in text if c in ('\t', '\n', '\r') or (ord(c) >= 32 and ord(c) != 127)]
    return ''.join(clean_chars)

lines = md_content.split('\n')
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line or line.startswith('# BÁO CÁO') or line.startswith('**Tuần 7') or line.startswith('**Đề tài') or line.startswith('**Người thực hiện') or line == '---':
        i += 1
        continue
    
    if line.startswith('## '):
        p = doc.add_paragraph()
        p.add_run(sanitize_text(line.replace('## ', '')))
        set_para_font(p, size=14, color=COLOR_PRIMARY, bold=True, space_after=8)
        i += 1
    elif line.startswith('### '):
        p = doc.add_paragraph()
        p.add_run(sanitize_text(line.replace('### ', '')))
        set_para_font(p, size=12.5, color=COLOR_SECONDARY, bold=True, space_after=6)
        i += 1
    elif line.startswith('#### '):
        p = doc.add_paragraph()
        p.add_run(sanitize_text(line.replace('#### ', '')))
        set_para_font(p, size=11.5, color=COLOR_PRIMARY, bold=True, italic=True, space_after=4)
        i += 1
    elif line.startswith('|') and '---' not in line:
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            if '---|---' not in lines[i]:
                table_lines.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        if table_lines:
            table = doc.add_table(rows=len(table_lines), cols=len(table_lines[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate(table_lines):
                for c_idx, cell_text in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    clean_cell = sanitize_text(cell_text)
                    cell.text = clean_cell
                    p_c = cell.paragraphs[0]
                    if r_idx == 0:
                        set_para_font(p_c, size=9.5, color=RGBColor(255, 255, 255), bold=True, space_after=2)
                        shd = OxmlElement('w:shd')
                        shd.set(qn('w:val'), 'clear')
                        shd.set(qn('w:color'), 'auto')
                        shd.set(qn('w:fill'), '1A365D')
                        cell._tc.get_or_add_tcPr().append(shd)
                    else:
                        set_para_font(p_c, size=9.5, color=COLOR_TEXT, bold=('**' in cell_text), space_after=2)
                        if r_idx % 2 == 1:
                            shd = OxmlElement('w:shd')
                            shd.set(qn('w:val'), 'clear')
                            shd.set(qn('w:color'), 'auto')
                            shd.set(qn('w:fill'), 'F7FAFC')
                            cell._tc.get_or_add_tcPr().append(shd)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
    else:
        p = doc.add_paragraph()
        p.add_run(sanitize_text(line))
        set_para_font(p, size=11, color=COLOR_TEXT, space_after=4)
        i += 1

doc.save(docx_path)
print('Successfully saved DOCX to', docx_path)
